from __future__ import annotations

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pytest import MonkeyPatch

from app.api import routes
from app.db.session import SessionLocal
from app.models import AgentLog, Candidate, Job, JobStatus, Resume, ResumeStatus, Score
from app.services import agent_runner


def _resume_document() -> bytes:
    """Create a small valid DOCX resume for deterministic Mock-mode tests."""
    document = Document()
    document.add_paragraph("Name: Alice Example")
    document.add_paragraph("Skills: Python, FastAPI, PostgreSQL, Docker")
    document.add_paragraph("Projects: Built a FastAPI service with PostgreSQL.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _create_job(client: TestClient) -> int:
    """Create a reusable job and return its API identifier."""
    response = client.post(
        "/api/jobs",
        json={
            "title": "Backend Engineer",
            "jd_text": "Need Python, FastAPI, PostgreSQL, Docker and project experience.",
        },
    )
    assert response.status_code == 201
    return response.json()["id"]


def _upload_resume(client: TestClient, job_id: int, content: bytes, filename: str = "alice.docx") -> None:
    """Upload one valid DOCX resume to an existing job."""
    response = client.post(
        f"/api/jobs/{job_id}/resumes",
        files={"files": (filename, content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200


def _error_payload(response: object) -> None:
    """Assert the public API error contract and correlation header."""
    payload = response.json()
    assert set(("code", "message", "request_id")).issubset(payload)
    assert payload["request_id"]
    assert response.headers["X-Request-ID"] == payload["request_id"]


def test_create_job(client: TestClient) -> None:
    """Create a job through the public API."""
    request_id = "normal-success-request"
    response = client.post(
        "/api/jobs",
        json={"title": "Backend Engineer", "jd_text": "Need Python."},
        headers={"X-Request-ID": request_id},
    )
    assert response.status_code == 201
    assert response.headers["X-Request-ID"] == request_id
    job_id = response.json()["id"]
    response = client.get(f"/api/jobs/{job_id}")
    assert response.status_code == 200
    assert response.json()["status"] == "draft"


def test_upload_rejects_unsupported_format(client: TestClient) -> None:
    """Reject unsupported uploads with the stable public error structure."""
    job_id = _create_job(client)
    response = client.post(f"/api/jobs/{job_id}/resumes", files={"files": ("resume.txt", b"not a resume", "text/plain")})
    assert response.status_code == 400
    assert response.json()["code"] == "invalid_file_type"
    _error_payload(response)


def test_mock_screening_flow(client: TestClient) -> None:
    """Run the complete upload-to-ranking path in deterministic Mock mode."""
    job_id = _create_job(client)
    _upload_resume(client, job_id, _resume_document())

    response = client.post(f"/api/jobs/{job_id}/run")
    assert response.status_code == 200
    assert response.json()["status"] == "running"

    job_response = client.get(f"/api/jobs/{job_id}")
    results_response = client.get(f"/api/jobs/{job_id}/results")
    logs_response = client.get(f"/api/jobs/{job_id}/logs")
    assert job_response.json()["status"] == "completed"
    assert len(results_response.json()) == 1
    assert results_response.json()[0]["score"]["total_score"] > 0
    assert any(log["step"] == "job_completed" for log in logs_response.json())


def test_rerun_reuses_single_candidate_and_score(client: TestClient) -> None:
    """Safely rerun completed and failed jobs without duplicate derived records or logs."""
    job_id = _create_job(client)
    document = _resume_document()
    _upload_resume(client, job_id, document)
    assert client.post(f"/api/jobs/{job_id}/run").status_code == 200

    db = SessionLocal()
    try:
        resume = db.query(Job).filter(Job.id == job_id).one().resumes[0]
        candidate = db.query(Candidate).filter(Candidate.resume_id == resume.id).one()
        candidate_id = candidate.id
        score_id = candidate.score.id
        resume_path = resume.file_path
        assert db.query(Candidate).filter(Candidate.resume_id == resume.id).count() == 1
        assert db.query(Score).filter(Score.candidate_id == candidate.id).count() == 1
    finally:
        db.close()

    assert client.post(f"/api/jobs/{job_id}/run").status_code == 200
    db = SessionLocal()
    try:
        candidate = db.query(Candidate).filter(Candidate.resume_id == resume.id).one()
        assert candidate.id == candidate_id
        assert candidate.score.id == score_id
        assert db.query(Candidate).filter(Candidate.resume_id == resume.id).count() == 1
        assert db.query(Score).filter(Score.candidate_id == candidate.id).count() == 1
        assert db.query(AgentLog).filter(AgentLog.job_id == job_id, AgentLog.step == "job_started").count() == 1
        assert db.query(AgentLog).filter(AgentLog.job_id == job_id, AgentLog.step == "job_completed").count() == 1
    finally:
        db.close()

    with open(resume_path, "wb") as resume_file:
        resume_file.write(b"invalid docx")
    assert client.post(f"/api/jobs/{job_id}/run").status_code == 200

    db = SessionLocal()
    try:
        assert db.query(Job).filter(Job.id == job_id).one().status == JobStatus.failed
        assert db.query(Candidate).filter(Candidate.resume_id == resume.id).count() == 0
        assert db.query(Score).filter(Score.job_id == job_id).count() == 0
    finally:
        db.close()

    with open(resume_path, "wb") as resume_file:
        resume_file.write(document)
    assert client.post(f"/api/jobs/{job_id}/run").status_code == 200
    db = SessionLocal()
    try:
        assert db.query(Job).filter(Job.id == job_id).one().status == JobStatus.completed
        assert db.query(Candidate).filter(Candidate.resume_id == resume.id).count() == 1
        assert db.query(Score).filter(Score.job_id == job_id).count() == 1
    finally:
        db.close()


def test_running_job_rejects_a_second_run(client: TestClient) -> None:
    """Reject a second start request while a job is already running."""
    job_id = _create_job(client)
    _upload_resume(client, job_id, _resume_document())
    db = SessionLocal()
    try:
        job = db.query(Job).filter(Job.id == job_id).one()
        job.status = JobStatus.running
        db.commit()
    finally:
        db.close()

    request_id = "already-running-request"
    response = client.post(f"/api/jobs/{job_id}/run", headers={"X-Request-ID": request_id})
    assert response.status_code == 409
    assert response.json()["code"] == "job_running"
    assert response.json()["request_id"] == request_id
    assert response.headers["X-Request-ID"] == request_id


def test_one_failed_resume_preserves_other_completed_results(client: TestClient) -> None:
    """Keep a completed resume result when another resume fails in the same job run."""
    job_id = _create_job(client)
    _upload_resume(client, job_id, _resume_document())
    _upload_resume(client, job_id, b"invalid docx", filename="broken.docx")

    assert client.post(f"/api/jobs/{job_id}/run").status_code == 200
    db = SessionLocal()
    try:
        resumes = {resume.original_filename: resume for resume in db.query(Resume).filter(Resume.job_id == job_id).all()}
        completed_resume = resumes["alice.docx"]
        failed_resume = next(resume for name, resume in resumes.items() if name != "alice.docx")
        assert completed_resume.status == ResumeStatus.completed
        assert failed_resume.status == ResumeStatus.failed
        assert db.query(Candidate).filter(Candidate.resume_id == completed_resume.id).count() == 1
        assert db.query(Candidate).filter(Candidate.resume_id == failed_resume.id).count() == 0
        assert db.query(Score).filter(Score.job_id == job_id).count() == 1
        assert db.query(Job).filter(Job.id == job_id).one().status == JobStatus.completed
    finally:
        db.close()


def test_outer_background_failure_converges_job_and_unfinished_resumes(client: TestClient, monkeypatch: MonkeyPatch) -> None:
    """Converge a job after an exception escapes the top-level background runner."""
    job_id = _create_job(client)
    _upload_resume(client, job_id, _resume_document())
    _upload_resume(client, job_id, b"invalid docx", filename="broken.docx")
    assert client.post(f"/api/jobs/{job_id}/run").status_code == 200

    db = SessionLocal()
    try:
        resumes = db.query(Resume).filter(Resume.job_id == job_id).order_by(Resume.id).all()
        completed_resume, interrupted_resume = resumes
        assert completed_resume.status == ResumeStatus.completed
        interrupted_resume.status = ResumeStatus.uploaded
        job = db.query(Job).filter(Job.id == job_id).one()
        job.status = JobStatus.running
        db.commit()
    finally:
        db.close()

    def raise_unexpected_error(*args: object, **kwargs: object) -> None:
        """Simulate an unexpected failure that escapes the regular runner path."""
        raise RuntimeError("/private/path api-key Traceback must not persist")

    monkeypatch.setattr(agent_runner, "_run", raise_unexpected_error)
    agent_runner.run_screening_job(job_id)

    db = SessionLocal()
    try:
        job = db.query(Job).filter(Job.id == job_id).one()
        resumes = db.query(Resume).filter(Resume.job_id == job_id).order_by(Resume.id).all()
        assert job.status == JobStatus.failed
        assert job.error_message == agent_runner.SAFE_JOB_FAILURE_MESSAGE
        assert resumes[0].status == ResumeStatus.completed
        assert resumes[1].status == ResumeStatus.failed
        assert resumes[1].error_message == agent_runner.SAFE_RESUME_FAILURE_MESSAGE
        assert db.query(Candidate).filter(Candidate.resume_id == resumes[0].id).count() == 1
    finally:
        db.close()


def test_error_responses_are_safe_and_consistent(client: TestClient, monkeypatch: MonkeyPatch) -> None:
    """Distinguish 404, validation, and business failures with the same error shape."""
    request_id = "client-error-request"
    not_found = client.get("/api/jobs/999999", headers={"X-Request-ID": request_id})
    assert not_found.status_code == 404
    assert not_found.json()["code"] == "job_not_found"
    assert not_found.json()["request_id"] == request_id
    _error_payload(not_found)

    validation = client.post("/api/jobs", json={"title": "", "jd_text": ""})
    assert validation.status_code == 422
    assert validation.json()["code"] == "validation_error"
    _error_payload(validation)

    job_id = _create_job(client)
    no_resumes = client.post(f"/api/jobs/{job_id}/run")
    assert no_resumes.status_code == 400
    assert no_resumes.json()["code"] == "no_resumes"
    _error_payload(no_resumes)

    def raise_unexpected_error(*args: object, **kwargs: object) -> object:
        """Simulate an implementation failure without exposing its content to clients."""
        raise RuntimeError("database URL /secret/path must not be exposed")

    monkeypatch.setattr(routes, "_get_job_or_404", raise_unexpected_error)
    internal_error = client.get("/api/jobs/1")
    assert internal_error.status_code == 500
    assert internal_error.json()["code"] == "internal_error"
    error_message = internal_error.json()["message"].lower()
    for unsafe_value in ("secret", "private", "path", "traceback", "runtimeerror"):
        assert unsafe_value not in error_message
    _error_payload(internal_error)
